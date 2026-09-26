"""Local-only extraction and provenance indexing for recovered mail attachments.

Raw attachments and extracted text stay under .agent-local or the user's
configured attachment folder. Nothing in this module promotes attachment
contents into Git-backed profile state.
"""
from __future__ import annotations

import html
import json
import re
import sqlite3
import zipfile
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from . import config
from ..privacy import privacy_findings

SCHEMA = "agent-os/mail-attachment-analysis/v1"
MAX_TEXT = 2_000_000
MAX_SHEET_CELLS = 50_000
TEXTISH_BYTES = set(range(32, 127)) | {9, 10, 13}


def _sniff_kind(path: Path, mime_type: str) -> str:
    ext = path.suffix.casefold()
    if ext:
        return ext
    try:
        head = path.read_bytes()[:8192]
    except OSError:
        return ""
    if head.startswith(b"%PDF-"):
        return ".pdf"
    if head.startswith((b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff")):
        return ".image"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                if "word/document.xml" in names:
                    return ".docx"
                if "xl/workbook.xml" in names:
                    return ".xlsx"
                if "META-INF/container.xml" in names:
                    return ".epub"
        except (OSError, zipfile.BadZipFile):
            pass
    if mime_type in {"text/calendar", "application/ics"}:
        return ".ics"
    if head:
        ratio = sum(byte in TEXTISH_BYTES for byte in head) / len(head)
        if ratio > 0.92:
            return ".txt"
    return ""


class _TextHTML(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data)

    def text(self) -> str:
        return html.unescape("\n".join(self.parts))


def _trim(text: str) -> str:
    return text[:MAX_TEXT]


def _plain(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return _trim(raw.decode(encoding)), f"text:{encoding}"
        except UnicodeDecodeError:
            pass
    return "", "text:decode_failed"


def _html_text(raw: str) -> str:
    parser = _TextHTML()
    parser.feed(raw)
    return parser.text()


def _pdf_ocr(path: Path) -> tuple[str, str]:
    try:
        import pymupdf
        from PIL import Image
        import pytesseract
    except ImportError:
        return "", "pdf-ocr:unavailable"
    chunks: list[str] = []
    try:
        with pymupdf.open(path) as document:
            for page in document:
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                chunks.append(pytesseract.image_to_string(image))
                if sum(map(len, chunks)) >= MAX_TEXT:
                    break
    except Exception:
        return "", "pdf-ocr:failed"
    return _trim("\n".join(chunks)), "pymupdf+pytesseract"


def _pdf(path: Path, *, ocr: bool) -> tuple[str, str]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted and not reader.decrypt(""):
        return "", "pypdf:encrypted"
    text = _trim("\n".join((page.extract_text() or "") for page in reader.pages))
    if text.strip():
        return text, "pypdf"
    if ocr:
        return _pdf_ocr(path)
    return "", "pypdf:needs_ocr"


def _docx(path: Path) -> tuple[str, str]:
    from docx import Document
    document = Document(str(path))
    chunks = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))
    return _trim("\n".join(chunks)), "python-docx"


def _xlsx(path: Path) -> tuple[str, str]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    chunks: list[str] = []
    cells = 0
    try:
        for sheet in workbook.worksheets:
            chunks.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    chunks.append("\t".join(values))
                cells += len(values)
                if cells >= MAX_SHEET_CELLS:
                    chunks.append("[TRUNCATED: cell limit]")
                    return _trim("\n".join(chunks)), "openpyxl:cell_limit"
    finally:
        workbook.close()
    return _trim("\n".join(chunks)), "openpyxl"


def _epub(path: Path) -> tuple[str, str]:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if Path(name).suffix.casefold() not in {".html", ".htm", ".xhtml"}:
                continue
            try:
                raw = archive.read(name).decode("utf-8", errors="replace")
            except (KeyError, OSError):
                continue
            chunks.append(_html_text(raw))
            if sum(map(len, chunks)) >= MAX_TEXT:
                break
    return _trim("\n".join(chunks)), "epub-html"


def _ics(path: Path) -> tuple[str, str]:
    raw, _ = _plain(path)
    unfolded = re.sub(r"\r?\n[ \t]", "", raw)
    keep = {
        "SUMMARY", "DTSTART", "DTEND", "DUE", "LOCATION", "DESCRIPTION",
        "ORGANIZER", "ATTENDEE", "URL", "STATUS", "TRANSP",
    }
    rows: list[str] = []
    for line in unfolded.splitlines():
        field = line.split(":", 1)[0].split(";", 1)[0].upper()
        if field in keep:
            rows.append(line)
    return _trim("\n".join(rows)), "icalendar-text"


def _media(path: Path) -> tuple[str, str]:
    try:
        import mutagen
    except ImportError:
        return "", "media:dependency_missing"
    try:
        media = mutagen.File(path)
    except Exception:
        return "", "media:parse_failed"
    if media is None:
        return "", "media:unsupported"
    rows: list[str] = []
    info = getattr(media, "info", None)
    length = getattr(info, "length", None)
    if isinstance(length, (int, float)):
        rows.append(f"duration_seconds: {length:.3f}")
    tags = getattr(media, "tags", None)
    if tags:
        for key in sorted(tags.keys(), key=str):
            value = tags.get(key)
            rendered = str(value)
            if rendered and len(rendered) <= 2000:
                rows.append(f"{key}: {rendered}")
    return _trim("\n".join(rows)), "mutagen"


def _image(path: Path, *, ocr: bool) -> tuple[str, str]:
    if not ocr:
        return "", "image:needs_ocr"
    try:
        from PIL import Image
        import pytesseract
        try:
            from pillow_heif import register_heif_opener
            register_heif_opener()
        except ImportError:
            pass
    except ImportError:
        return "", "image:ocr_unavailable"
    try:
        with Image.open(path) as image:
            text = pytesseract.image_to_string(image)
        return _trim(text), "pytesseract"
    except Exception:
        return "", "image:ocr_failed"


def extract_text(path: Path, *, mime_type: str = "", ocr: bool = False) -> dict[str, Any]:
    ext = _sniff_kind(path, mime_type) or path.suffix.casefold()
    try:
        if ext in {".txt", ".csv", ".md", ".json", ".xml"} or mime_type.startswith("text/plain"):
            text, method = _plain(path)
        elif ext in {".html", ".htm"} or mime_type == "text/html":
            raw, _ = _plain(path)
            text, method = _html_text(raw), "html-parser"
        elif ext == ".ics" or mime_type in {"text/calendar", "application/ics"}:
            text, method = _ics(path)
        elif ext == ".pdf" or mime_type == "application/pdf":
            text, method = _pdf(path, ocr=ocr)
        elif ext == ".docx" or "wordprocessingml" in mime_type:
            text, method = _docx(path)
        elif ext == ".xlsx" or "spreadsheetml" in mime_type:
            text, method = _xlsx(path)
        elif ext == ".epub" or mime_type == "application/epub+zip":
            text, method = _epub(path)
        elif ext in {".png", ".jpg", ".jpeg", ".heic", ".webp", ".image"} or mime_type.startswith("image/"):
            text, method = _image(path, ocr=ocr)
        elif ext in {".m4r", ".m4a", ".mp4", ".mov"} or mime_type.startswith(("audio/", "video/")):
            text, method = _media(path)
        else:
            return {"status": "unsupported", "method": "metadata-only", "text": ""}
    except ImportError as exc:
        return {"status": "dependency_missing", "method": type(exc).__name__, "text": ""}
    except Exception as exc:
        return {"status": "error", "method": type(exc).__name__, "text": ""}
    needs_ocr = method in {"image:needs_ocr", "pypdf:needs_ocr"}
    unavailable_ocr = method in {"image:ocr_unavailable", "pdf-ocr:unavailable"}
    failed = method in {"image:ocr_failed", "pdf-ocr:failed", "media:parse_failed"}
    dep_missing = method == "media:dependency_missing"
    status = (
        "parsed" if text.strip()
        else "needs_ocr" if needs_ocr
        else "ocr_unavailable" if unavailable_ocr
        else "dependency_missing" if dep_missing
        else "error" if failed
        else "empty_text"
    )
    return {"status": status, "method": method, "text": text}


def _safe_name(value: str) -> str:
    return re.sub(r"[^a-zA-Z0-9._-]+", "_", value)[:120] or "attachment"


def classify_for_review(filename: str, mime_type: str, text: str = "") -> dict[str, Any]:
    """Local-only triage for profile/relationship review.

    This deliberately classifies *what deserves review*, not what is true
    about the user. The returned category never implies promotion.
    """
    name = (filename or "").casefold()
    sample = (text or "")[:50000].casefold()
    haystack = f"{name}\n{sample}"
    ext = Path(filename or "").suffix.casefold()

    def any_term(*terms: str) -> bool:
        return any(term in haystack for term in terms)

    category = "general_document"
    priority = "low"
    sensitivity = "personal_private"
    destinations: list[str] = []

    if any_term("resume", "curriculum vitae") or ext in {".doc", ".docx"} and any_term("education", "employment", "experience", "skills"):
        category, priority = "resume_or_cv", "high"
        sensitivity = "restricted_personal"
        destinations = ["education", "work_history", "skills"]
    elif any_term("lease", "landlord", "tenant", "rental application"):
        category, priority = "housing_or_lease", "high"
        sensitivity = "restricted_personal"
        destinations = ["housing_timeline", "household"]
    elif any_term("offer letter", "employment application", "employment verification", "w-2", "w2", "w-4", "w4", "payroll", "benefits enrollment"):
        category, priority = "employment_or_hr", "high"
        sensitivity = "restricted_personal"
        destinations = ["work_history", "benefits_or_finance"]
    elif any_term("transcript", "diploma", "degree", "university", "college", "student"):
        category, priority = "education", "high"
        sensitivity = "restricted_personal"
        destinations = ["education"]
    elif any_term("vehicle", "vin", "auto insurance", "insurance card", "registration"):
        category, priority = "vehicle_or_insurance", "medium"
        sensitivity = "restricted_personal"
        destinations = ["vehicle", "finance"]
    elif ext in {".xlsx", ".xls", ".csv"} or "spreadsheet" in mime_type.casefold():
        category, priority = "spreadsheet", "medium"
        sensitivity = "personal_private"
        destinations = ["planning", "finance_or_projects"]
    elif ext == ".epub" or mime_type == "application/epub+zip":
        category, priority = "ebook", "low"
        sensitivity = "personal_private"
        destinations = ["reading_interests"]
    elif ext in {".ics", ".vcs"} or mime_type in {"text/calendar", "application/ics"}:
        category, priority = "calendar", "medium"
        sensitivity = "personal_private"
        destinations = ["timeline", "commitments"]
    elif mime_type.startswith("image/") or ext in {".png", ".jpg", ".jpeg", ".heic", ".webp"}:
        category, priority = "image_or_scan", "medium"
        sensitivity = "personal_private"
        destinations = ["visual_evidence"]
    elif mime_type.startswith(("audio/", "video/")) or ext in {".m4r", ".m4a", ".mp4", ".mov"}:
        category, priority = "media", "low"
        sensitivity = "personal_private"
        destinations = ["media_or_routine"]

    findings = privacy_findings((text or "")[:50000])
    if any_term("social security", "ssn", "passport", "driver's license", "medical", "diagnosis", "bank account", "routing number") or findings:
        sensitivity = "restricted_high"
        priority = "high"

    return {
        "review_category": category,
        "review_priority": priority,
        "sensitivity_hint": sensitivity,
        "candidate_destinations": destinations,
        "privacy_findings": findings,
        "promotion_eligible": False,
    }


def analyze_recovered(
    scope: str,
    *,
    root: Path | None = None,
    ocr: bool = False,
    limit: int | None = None,
    manifest_path: Path | None = None,
    mail_db: Path | None = None,
    output_dir: Path | None = None,
) -> dict[str, Any]:
    """Extract recovered attachment text and attach it back to the local mail index.

    Each record preserves source_id -> message_id -> attachment_id provenance.
    Content remains local/private and is never a profile promotion by itself.
    """
    scope = config.check_scope(scope)
    manifest = Path(manifest_path or (config.scope_dir(scope) / "attachments.sqlite3"))
    db_path = Path(mail_db or config.db_for(scope))
    out = Path(output_dir or (config.scope_dir(scope) / "attachment-analysis"))
    text_dir = out / "text"
    out.mkdir(parents=True, exist_ok=True)
    text_dir.mkdir(parents=True, exist_ok=True)
    if not manifest.exists():
        raise FileNotFoundError(f"attachment export manifest not found: {manifest}")
    if not db_path.exists():
        raise FileNotFoundError(f"mail index not found: {db_path}")

    records: list[dict[str, Any]] = []
    counts: Counter[str] = Counter()
    with sqlite3.connect(manifest) as mconn, sqlite3.connect(db_path) as conn:
        mconn.row_factory = sqlite3.Row
        conn.row_factory = sqlite3.Row
        try:
            occurrence_rows = mconn.execute(
                "SELECT o.source_id,o.sha256,f.path FROM occurrences o JOIN files f ON f.sha256=o.sha256 "
                "UNION SELECT source_id,sha256,path FROM files ORDER BY path,source_id"
            ).fetchall()
        except sqlite3.OperationalError:
            occurrence_rows = mconn.execute("SELECT source_id,sha256,path FROM files ORDER BY path,source_id").fetchall()
        if limit is not None:
            occurrence_rows = occurrence_rows[: max(0, limit)]
        touched_messages: set[str] = set()
        processed_attachment_ids: set[str] = set()
        for row in occurrence_rows:
            path = Path(row["path"])
            source = conn.execute(
                "SELECT s.message_id,m.subject,m.sender,m.date_utc,s.source_kind,s.account_id,s.mailbox "
                "FROM mail_sources s JOIN mail_messages m ON m.id=s.message_id WHERE s.id=?",
                (row["source_id"],),
            ).fetchone()
            if source is None:
                counts["source_unresolved"] += 1
                continue
            attachments = conn.execute(
                "SELECT id,filename,mime_type,size FROM mail_attachments WHERE message_id=? AND content_hash=? ORDER BY id",
                (source["message_id"], row["sha256"]),
            ).fetchall()
            exempl = attachments[0] if attachments else None
            mime = str(exempl["mime_type"] or "") if exempl else ""
            result = extract_text(path, mime_type=mime, ocr=ocr) if path.exists() else {"status":"missing_file","method":"none","text":""}
            text = result.pop("text", "")
            text_path = None
            if text:
                text_path = text_dir / f"{_safe_name(str(row['sha256']))}.txt"
                text_path.write_text(text, encoding="utf-8")
            targets = attachments or [None]
            for attachment in targets:
                review = classify_for_review(attachment["filename"] if attachment else path.name, str(attachment["mime_type"] or "") if attachment else mime, text)
                record = {
                    "schema": SCHEMA,
                    "scope": scope,
                    "source_id": row["source_id"],
                    "message_id": source["message_id"],
                    "attachment_id": attachment["id"] if attachment else None,
                    "sha256": row["sha256"],
                    "path": str(path),
                    "filename": attachment["filename"] if attachment else path.name,
                    "mime_type": str(attachment["mime_type"] or "") if attachment else mime,
                    "message_subject": source["subject"],
                    "message_sender": source["sender"],
                    "message_date": source["date_utc"],
                    "source_kind": source["source_kind"],
                    "account_id": source["account_id"],
                    "mailbox": source["mailbox"],
                    "privacy": "local_private",
                    **review,
                    **result,
                    "text_chars": len(text),
                }
                if text_path is not None:
                    record["text_path"] = str(text_path)
                records.append(record)
                counts[result["status"]] += 1
                if attachment:
                    processed_attachment_ids.add(str(attachment["id"]))
                if attachment and result["status"] in {"parsed", "empty_text", "needs_ocr", "ocr_unavailable", "dependency_missing", "error"}:
                    conn.execute("UPDATE mail_attachments SET text=?, extraction=? WHERE id=?", (text or None, result["method"], attachment["id"]))
                    touched_messages.add(str(source["message_id"]))

        indexed_rows = conn.execute(
            "SELECT a.id,a.message_id,a.filename,a.mime_type,a.size,a.content_hash,"
            "m.subject,m.sender,m.date_utc FROM mail_attachments a "
            "JOIN mail_messages m ON m.id=a.message_id ORDER BY a.message_id,a.id"
        ).fetchall()
        for attachment in indexed_rows:
            if str(attachment["id"]) in processed_attachment_ids:
                continue
            filename = str(attachment["filename"] or "")
            mime = str(attachment["mime_type"] or "")
            ext = Path(filename).suffix.casefold()
            intentional = mime in {"application/pkcs7-signature", "application/x-pkcs7-signature"} or ext == ".p7s"
            status = "intentionally_skipped" if intentional else "not_recovered"
            method = "signature_metadata_only" if intentional else "indexed_metadata_only"
            review = classify_for_review(filename, mime, "")
            records.append({
                "schema": SCHEMA,
                "scope": scope,
                "source_id": None,
                "message_id": attachment["message_id"],
                "attachment_id": attachment["id"],
                "sha256": attachment["content_hash"],
                "path": None,
                "filename": attachment["filename"],
                "mime_type": attachment["mime_type"],
                "message_subject": attachment["subject"],
                "message_sender": attachment["sender"],
                "message_date": attachment["date_utc"],
                "privacy": "local_private",
                **review,
                "status": status,
                "method": method,
                "text_chars": 0,
            })
            counts[status] += 1

        for message_id in touched_messages:
            try:
                joined = conn.execute(
                    "SELECT GROUP_CONCAT(COALESCE(filename,'') || ' ' || COALESCE(text,''), char(10)) AS value "
                    "FROM mail_attachments WHERE message_id=?", (message_id,)
                ).fetchone()["value"] or ""
                conn.execute("UPDATE mail_fts SET attachments=? WHERE message_id=?", (joined, message_id))
            except sqlite3.OperationalError:
                break
        conn.commit()

    records_path = out / "records.jsonl"
    records_path.write_text("".join(json.dumps(item, ensure_ascii=False) + "\n" for item in records), encoding="utf-8")
    attachment_ids = {str(item["attachment_id"]) for item in records if item.get("attachment_id")}
    status_by_attachment: dict[str, set[str]] = {}
    for item in records:
        attachment_id = item.get("attachment_id")
        if attachment_id:
            status_by_attachment.setdefault(str(attachment_id), set()).add(str(item.get("status") or ""))

    gap_statuses = {"not_recovered", "missing_file"}
    pending_analysis_statuses = {"needs_ocr", "ocr_unavailable", "dependency_missing", "unsupported", "metadata_only", "error"}
    recovered_ids = {
        attachment_id
        for attachment_id, statuses in status_by_attachment.items()
        if statuses and not statuses.issubset(gap_statuses | {"intentionally_skipped"})
        and not statuses.intersection(gap_statuses)
    }
    intentionally_skipped_ids = {
        attachment_id
        for attachment_id, statuses in status_by_attachment.items()
        if statuses == {"intentionally_skipped"}
    }
    gap_ids = {
        attachment_id
        for attachment_id, statuses in status_by_attachment.items()
        if statuses.intersection(gap_statuses)
    }
    pending_ids = {
        attachment_id
        for attachment_id, statuses in status_by_attachment.items()
        if statuses.intersection(pending_analysis_statuses)
    }

    summary = {
        "schema": SCHEMA,
        "scope": scope,
        "attachment_records_accounted_for": len(attachment_ids),
        "attachment_records_recovered": len(recovered_ids),
        "attachment_records_intentionally_skipped": len(intentionally_skipped_ids),
        "attachment_records_not_recovered": len(gap_ids),
        "attachment_records_pending_analysis": len(pending_ids),
        "coverage_complete": bool(attachment_ids) and not gap_ids,
        "analysis_complete": bool(attachment_ids) and not gap_ids and not pending_ids,
        "occurrence_records_written": len(records),
        "counts": dict(sorted(counts.items())),
        "ocr_requested": bool(ocr),
        "records": str(records_path),
        "text_dir": str(text_dir),
        "privacy": "All binaries and extracted text remain local/private; no automatic profile promotion.",
    }
    review_rank = {"high": 0, "medium": 1, "low": 2}
    review_by_id: dict[str, dict[str, Any]] = {}
    for item in records:
        attachment_id = item.get("attachment_id")
        if not attachment_id:
            continue
        key = str(attachment_id)
        current = review_by_id.get(key)
        if current is None:
            current = {
                "attachment_id": attachment_id,
                "message_id": item.get("message_id"),
                "source_ids": [],
                "occurrence_count": 0,
                "filename": item.get("filename"),
                "mime_type": item.get("mime_type"),
                "statuses": [],
                "review_category": item.get("review_category"),
                "review_priority": item.get("review_priority"),
                "sensitivity_hint": item.get("sensitivity_hint"),
                "candidate_destinations": item.get("candidate_destinations", []),
                "privacy_findings": item.get("privacy_findings", []),
                "promotion_eligible": False,
            }
            review_by_id[key] = current
        current["occurrence_count"] += 1
        source_id = item.get("source_id")
        if source_id and source_id not in current["source_ids"]:
            current["source_ids"].append(source_id)
        status = str(item.get("status") or "")
        if status and status not in current["statuses"]:
            current["statuses"].append(status)
        current["privacy_findings"] = sorted(set(current["privacy_findings"]) | set(item.get("privacy_findings", [])))
        destinations = list(current["candidate_destinations"])
        for destination in item.get("candidate_destinations", []):
            if destination not in destinations:
                destinations.append(destination)
        current["candidate_destinations"] = destinations
        if review_rank.get(str(item.get("review_priority")), 9) < review_rank.get(str(current.get("review_priority")), 9):
            current["review_priority"] = item.get("review_priority")
            current["review_category"] = item.get("review_category")
            current["sensitivity_hint"] = item.get("sensitivity_hint")

    review_items = sorted(
        review_by_id.values(),
        key=lambda item: (review_rank.get(str(item.get("review_priority")), 9), str(item.get("filename") or "")),
    )
    review_path = out / "review-queue.jsonl"
    review_path.write_text("".join(json.dumps(item, ensure_ascii=False) + "\n" for item in review_items), encoding="utf-8")
    summary["review_queue"] = str(review_path)
    summary["review_priority_counts"] = dict(Counter(str(item.get("review_priority")) for item in review_items))
    summary["review_category_counts"] = dict(Counter(str(item.get("review_category")) for item in review_items))
    (out / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return summary
