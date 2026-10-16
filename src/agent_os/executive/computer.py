"""Tiered, path-guarded access to the operator's computer for the Executive.

Tiers (see docs/orchestrator-os/executive-window.md, "Computer access"):

- free: list_dir, read_text, search_files
- automatic and logged: write_draft (new files under drafts_root only)
- window Confirm: write_file, move_file, delete_file, open_path, run_command

This module does the work; it does not decide who may call what. The Confirm
tier is enforced by actions_computer.py (consequential window actions) and the
MCP tools, which only park a pending confirmation. There is no shell: commands
are argv lists run without one.

Every call appends {ts, op, path, ok} to <window home>/computer-audit.jsonl.
File contents are never written to the audit.
"""
from __future__ import annotations

import difflib
import fnmatch
import hashlib
import json
import os
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from ..paths import ROOT
from .window.confirm import window_home

ENV_CONFIG = "AGENT_OS_COMPUTER_CONFIG"
DEFAULT_CONFIG_PATH = Path.home() / ".agent-os" / "computer.json"
DEFAULT_DRAFTS_ROOT = ROOT / ".agent-local" / "executive" / "drafts"
DEFAULT_TRASH_ROOT = ROOT / ".agent-local" / "executive" / "trash"
AUDIT_FILE = "computer-audit.jsonl"

MAX_READ_BYTES = 2 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
MAX_LIST_ENTRIES = 500
MAX_SEARCH_RESULTS = 100
MAX_SEARCH_SCANNED = 20_000
MAX_DIFF_CHARS = 50_000
MAX_OUTPUT_CHARS = 20_000
COMMAND_TIMEOUT_SECONDS = 120

TEXT_SUFFIXES = frozenset({".txt", ".md", ".csv", ".json"})
DOC_SUFFIXES = frozenset({".docx", ".xlsx", ".pdf"})
# Characters a shell would interpret. No shell runs, but an argv carrying them
# is almost always a shell line pasted into a list, so it is refused outright.
SHELL_METACHARACTERS = frozenset(";&|<>`$\n\r")
SHELL_PROGRAMS = frozenset({
    "sh", "bash", "zsh", "dash", "ksh", "fish", "csh", "tcsh",
    "cmd", "powershell", "pwsh", "wsl", "bash.exe", "cmd.exe", "powershell.exe", "pwsh.exe", "wsl.exe",
})

TIERS = {
    "list_dir": "free",
    "read_text": "free",
    "search_files": "free",
    "write_draft": "automatic",
    "write_file": "confirm",
    "move_file": "confirm",
    "delete_file": "confirm",
    "open_path": "confirm",
    "run_command": "confirm",
}

# Injection hooks (tests): OPENER(path) replaces os.startfile / xdg-open; RUNNER replaces subprocess.run.
OPENER: Callable[[str], Any] | None = None
RUNNER: Callable[..., Any] | None = None


class ComputerAccessError(ValueError):
    """A request the guard refuses (outside the allowed roots, wrong shape, ...)."""


# ---------------------------------------------------------------- config


def config_path() -> Path:
    env = os.environ.get(ENV_CONFIG)
    return Path(env) if env else DEFAULT_CONFIG_PATH


def _default_roots() -> list[Path]:
    home = Path.home()
    return [home / name for name in ("Documents", "Desktop", "Downloads") if (home / name).is_dir()]


def load_config(path: Path | str | None = None) -> dict[str, Any]:
    """{allowed_roots, drafts_root, trash_root} as resolved Paths; defaults fill gaps."""
    source = Path(path) if path is not None else config_path()
    try:
        data = json.loads(source.read_text(encoding="utf-8"))
    except FileNotFoundError:
        data = {}
    except (OSError, ValueError) as exc:
        raise ComputerAccessError(f"unreadable computer config {source}: {exc}") from exc
    if not isinstance(data, dict):
        raise ComputerAccessError(f"computer config {source} must be a JSON object")
    roots = data.get("allowed_roots")
    if roots is None:
        root_paths = _default_roots()
    elif isinstance(roots, list) and all(isinstance(r, str) and r.strip() for r in roots):
        root_paths = [Path(os.path.expandvars(r)).expanduser() for r in roots]
    else:
        raise ComputerAccessError("allowed_roots must be a list of paths")
    drafts = data.get("drafts_root")
    trash = data.get("trash_root")
    return {
        "allowed_roots": [p.resolve() for p in root_paths],
        "drafts_root": (Path(os.path.expandvars(drafts)).expanduser() if drafts else DEFAULT_DRAFTS_ROOT).resolve(),
        "trash_root": (Path(os.path.expandvars(trash)).expanduser() if trash else DEFAULT_TRASH_ROOT).resolve(),
    }


# ---------------------------------------------------------------- guard


def _within(path: Path, root: Path) -> bool:
    a, b = os.path.normcase(str(path)), os.path.normcase(str(root))
    try:
        return os.path.commonpath([a, b]) == b
    except ValueError:  # different drives on Windows
        return False


def guard(path: Any, cfg: dict[str, Any] | None = None, roots: list[Path] | None = None) -> Path:
    """Resolve symlinks and '..' and return the path only if it lies under an allowed root.

    Relative paths are taken against the first allowed root."""
    if not isinstance(path, (str, os.PathLike)) or not str(path).strip():
        raise ComputerAccessError("path must be a non-empty string")
    if "\x00" in str(path):
        raise ComputerAccessError("path contains a NUL byte")
    cfg = cfg or load_config()
    allowed = roots if roots is not None else cfg["allowed_roots"]
    if not allowed:
        raise ComputerAccessError("no allowed roots are configured")
    raw = Path(os.path.expandvars(str(path))).expanduser()
    if not raw.is_absolute():
        raw = allowed[0] / raw
    resolved = raw.resolve()
    if not any(_within(resolved, root) for root in allowed):
        raise ComputerAccessError(f"path is outside the allowed roots: {path}")
    return resolved


# ---------------------------------------------------------------- audit


def audit_path() -> Path:
    return window_home() / AUDIT_FILE


def _audit(op: str, path: Any, ok: bool, **extra: Any) -> None:
    line = {"ts": datetime.now(timezone.utc).isoformat(timespec="seconds"), "op": op,
            "path": str(path) if path is not None else None, "ok": bool(ok), **extra}
    try:
        target = audit_path()
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("a", encoding="utf-8", newline="\n") as handle:
            handle.write(json.dumps(line, sort_keys=True, default=str) + "\n")
    except OSError:
        pass  # the audit is best-effort; it must never break or alter the operation


def _audited(op: str, path: Any, fn: Callable[[], dict[str, Any]], **extra: Any) -> dict[str, Any]:
    try:
        result = fn()
    except Exception as exc:
        _audit(op, path, False, error=type(exc).__name__, **extra)
        raise
    _audit(op, result.get("path", path), True, **extra)
    return result


# ---------------------------------------------------------------- helpers


def _stat(path: Path) -> dict[str, Any]:
    st = path.stat()
    return {"path": str(path), "name": path.name, "type": "dir" if path.is_dir() else "file",
            "size": st.st_size if path.is_file() else None,
            "modified": datetime.fromtimestamp(st.st_mtime, timezone.utc).isoformat(timespec="seconds")}


def _sha256(path: Path) -> str | None:
    if not path.is_file():
        return None
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _extract(path: Path) -> str | None:
    """Text of a supported file, or None when its type (or optional library) is unsupported."""
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            return None
        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            return None
        book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            lines = []
            for sheet in book.worksheets:
                lines.append(f"# {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    lines.append(",".join("" if v is None else str(v) for v in row))
            return "\n".join(lines)
        finally:
            book.close()
    if suffix == ".pdf":
        try:
            import pypdf
        except ImportError:
            return None
        return "\n".join((page.extract_text() or "") for page in pypdf.PdfReader(str(path)).pages)
    return None


def _unique(path: Path, tag: str) -> Path:
    """path, else <stem>-<tag>N<suffix> for the first N >= 2 that does not exist."""
    if not path.exists():
        return path
    n = 2
    while True:
        candidate = path.with_name(f"{path.stem}-{tag}{n}{path.suffix}")
        if not candidate.exists():
            return candidate
        n += 1


def _cap(text: str, limit: int) -> tuple[str, bool]:
    return (text[:limit], True) if len(text) > limit else (text, False)


# ---------------------------------------------------------------- free tier


def list_dir(path: Any) -> dict[str, Any]:
    def run() -> dict[str, Any]:
        target = guard(path)
        if not target.is_dir():
            raise ComputerAccessError(f"not a directory: {path}")
        entries = []
        children = sorted(target.iterdir(), key=lambda p: p.name.lower())
        for child in children[:MAX_LIST_ENTRIES]:
            try:
                entries.append({k: v for k, v in _stat(child).items() if k != "path"})
            except OSError:
                entries.append({"name": child.name, "type": "unknown"})
        return {"path": str(target), "entries": entries, "truncated": len(children) > MAX_LIST_ENTRIES}

    return _audited("list_dir", path, run)


def read_text(path: Any) -> dict[str, Any]:
    def run() -> dict[str, Any]:
        target = guard(path)
        if not target.is_file():
            raise ComputerAccessError(f"not a file: {path}")
        meta = _stat(target)
        if meta["size"] > MAX_READ_BYTES:
            return {**meta, "text": None, "note": f"larger than {MAX_READ_BYTES} bytes; metadata only"}
        text = _extract(target)
        if text is None:
            return {**meta, "text": None, "note": "unsupported file type (or its reader is not installed); metadata only"}
        text, truncated = _cap(text, MAX_TEXT_CHARS)
        return {**meta, "text": text, "truncated": truncated}

    return _audited("read_text", path, run)


def search_files(root: Any, pattern: str = "*", contains: str | None = None, limit: int = MAX_SEARCH_RESULTS) -> dict[str, Any]:
    """Files under root whose name matches the glob (case-insensitive) and, if given, contain the substring."""
    def run() -> dict[str, Any]:
        cfg = load_config()
        base = guard(root, cfg)
        if not base.is_dir():
            raise ComputerAccessError(f"not a directory: {root}")
        if not isinstance(pattern, str) or not pattern:
            raise ComputerAccessError("pattern must be a non-empty glob")
        cap = max(1, min(int(limit), MAX_SEARCH_RESULTS))
        needle = contains.lower() if isinstance(contains, str) and contains else None
        glob = pattern.lower()
        results: list[dict[str, Any]] = []
        scanned = 0
        truncated = False
        for dirpath, dirnames, filenames in os.walk(base):  # never follows directory symlinks
            dirnames.sort()
            for name in sorted(filenames):
                scanned += 1
                if scanned > MAX_SEARCH_SCANNED or len(results) >= cap:
                    truncated = True
                    break
                if not fnmatch.fnmatchcase(name.lower(), glob):
                    continue
                try:
                    found = guard(Path(dirpath) / name, cfg)  # a file symlink may point outside
                except ComputerAccessError:
                    continue
                if needle is not None:
                    try:
                        if found.stat().st_size > MAX_READ_BYTES:
                            continue
                        text = _extract(found)
                    except Exception:
                        continue
                    if text is None or needle not in text.lower():
                        continue
                try:
                    results.append({k: v for k, v in _stat(found).items() if k != "type"})
                except OSError:
                    continue
            if truncated:
                break
        return {"path": str(base), "pattern": pattern, "contains": contains is not None,
                "results": results, "truncated": truncated}

    return _audited("search_files", root, run)


# ---------------------------------------------------------------- automatic tier


def write_draft(name: Any, content: Any) -> dict[str, Any]:
    """Create a new file under drafts_root. Never overwrites: an existing name gets -vN."""
    def run() -> dict[str, Any]:
        if not isinstance(name, str) or not name.strip() or name.strip() in (".", ".."):
            raise ComputerAccessError("name must be a plain file name")
        if any(sep in name for sep in ("/", "\\", ":")) or "\x00" in name:
            raise ComputerAccessError("name must be a plain file name, not a path")
        if not isinstance(content, str):
            raise ComputerAccessError("content must be text")
        drafts = load_config()["drafts_root"]
        drafts.mkdir(parents=True, exist_ok=True)
        target = guard(drafts / name.strip(), roots=[drafts])
        while True:
            target = _unique(target, "v")
            try:
                with target.open("x", encoding="utf-8", newline="") as handle:  # "x": never clobber
                    handle.write(content)
                break
            except FileExistsError:
                continue
        return {"path": str(target), "bytes": len(content.encode("utf-8"))}

    return _audited("write_draft", name, run)


# ---------------------------------------------------------------- confirm tier: previews


def preview_write(path: Any, content: Any) -> dict[str, Any]:
    def run() -> dict[str, Any]:
        if not isinstance(content, str):
            raise ComputerAccessError("content must be text")
        target = guard(path)
        if target.is_dir():
            raise ComputerAccessError(f"is a directory: {path}")
        exists = target.is_file()
        old = ""
        if exists:
            if target.stat().st_size > MAX_READ_BYTES:
                old = None
            else:
                try:
                    old = target.read_text(encoding="utf-8")
                except (UnicodeDecodeError, OSError):
                    old = None
        if old is None:
            diff = f"(existing file is binary or larger than {MAX_READ_BYTES} bytes; it will be replaced and backed up)"
        else:
            diff = "".join(difflib.unified_diff(
                old.splitlines(keepends=True), content.splitlines(keepends=True),
                fromfile=f"a/{target.name}" if exists else "/dev/null", tofile=f"b/{target.name}"))
        diff, truncated = _cap(diff, MAX_DIFF_CHARS)
        return {"op": "write_file", "path": str(target), "exists": exists, "backup": exists,
                "base_sha256": _sha256(target), "diff": diff, "diff_truncated": truncated}

    return _audited("preview_write", path, run)


def preview_move(source: Any, target: Any) -> dict[str, Any]:
    def run() -> dict[str, Any]:
        src, dst = _move_paths(source, target)
        return {"op": "move_file", "path": str(src), "source": str(src), "target": str(dst)}

    return _audited("preview_move", source, run)


def preview_delete(path: Any) -> dict[str, Any]:
    def run() -> dict[str, Any]:
        target = _delete_path(path)
        return {"op": "delete_file", "path": str(target), "trash_root": str(load_config()["trash_root"]),
                "note": "moved to the Agent OS trash, not deleted"}

    return _audited("preview_delete", path, run)


def preview_open(path: Any) -> dict[str, Any]:
    def run() -> dict[str, Any]:
        target = guard(path)
        if not target.exists():
            raise ComputerAccessError(f"does not exist: {path}")
        return {"op": "open_path", "path": str(target), "opener": _opener_name()}

    return _audited("preview_open", path, run)


def preview_run(argv: Any, cwd: Any = None) -> dict[str, Any]:
    def run() -> dict[str, Any]:
        args, workdir = _command(argv, cwd)
        return {"op": "run_command", "path": str(workdir), "argv": args, "cwd": str(workdir),
                "timeout_seconds": COMMAND_TIMEOUT_SECONDS, "shell": False}

    return _audited("preview_run", cwd, run)


# ---------------------------------------------------------------- confirm tier: execution


_UNSET: Any = object()


def write_file(path: Any, content: Any, expect_sha256: Any = _UNSET) -> dict[str, Any]:
    """Create or overwrite a file under the allowed roots, backing up an existing one to .bak first.

    expect_sha256 (the preview's base_sha256; None for a new file) refuses the
    write when the file changed since the operator saw the diff."""
    def run() -> dict[str, Any]:
        if not isinstance(content, str):
            raise ComputerAccessError("content must be text")
        target = guard(path)
        if target.is_dir():
            raise ComputerAccessError(f"is a directory: {path}")
        if expect_sha256 is not _UNSET and _sha256(target) != expect_sha256:
            raise ComputerAccessError("file changed since the preview; preview it again")
        backup = None
        if target.exists():
            backup = _unique(target.with_name(target.name + ".bak"), "")
            shutil.copy2(target, backup)
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("w", encoding="utf-8", newline="") as handle:
            handle.write(content)
        return {"path": str(target), "backup": str(backup) if backup else None,
                "bytes": len(content.encode("utf-8"))}

    return _audited("write_file", path, run)


def _move_paths(source: Any, target: Any) -> tuple[Path, Path]:
    cfg = load_config()
    src = guard(source, cfg)
    if not src.exists():
        raise ComputerAccessError(f"does not exist: {source}")
    if any(src == root for root in cfg["allowed_roots"]):
        raise ComputerAccessError("cannot move an allowed root")
    dst = guard(target, cfg)
    if dst.is_dir():
        dst = guard(dst / src.name, cfg)
    if dst.exists():
        raise ComputerAccessError(f"target already exists: {dst}")
    if _within(dst, src):
        raise ComputerAccessError("cannot move a directory into itself")
    return src, dst


def move_file(source: Any, target: Any) -> dict[str, Any]:
    """Move or rename within the allowed roots; never overwrites."""
    def run() -> dict[str, Any]:
        src, dst = _move_paths(source, target)
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(src), str(dst))
        return {"path": str(src), "source": str(src), "target": str(dst)}

    return _audited("move_file", source, run, target=str(target))


def _delete_path(path: Any) -> Path:
    cfg = load_config()
    target = guard(path, cfg)
    if not target.exists():
        raise ComputerAccessError(f"does not exist: {path}")
    if any(target == root for root in cfg["allowed_roots"]):
        raise ComputerAccessError("cannot delete an allowed root")
    return target


def delete_file(path: Any) -> dict[str, Any]:
    """Move to the Agent OS trash (with a sidecar naming the original path). Never hard-deletes."""
    def run() -> dict[str, Any]:
        target = _delete_path(path)
        trash = load_config()["trash_root"]
        trash.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        trashed = _unique(trash / f"{stamp}-{target.name}", "")
        shutil.move(str(target), str(trashed))
        sidecar = trashed.with_name(trashed.name + ".origin.json")
        sidecar.write_text(json.dumps({"original_path": str(target), "trashed_at": stamp}) + "\n", encoding="utf-8")
        return {"path": str(target), "trashed_to": str(trashed)}

    return _audited("delete_file", path, run)


def _opener_name() -> str:
    if sys.platform.startswith("win"):
        return "os.startfile"
    return "open" if sys.platform == "darwin" else "xdg-open"


def open_path(path: Any) -> dict[str, Any]:
    """Open a file or folder with its default application."""
    def run() -> dict[str, Any]:
        target = guard(path)
        if not target.exists():
            raise ComputerAccessError(f"does not exist: {path}")
        if OPENER is not None:
            OPENER(str(target))
        elif sys.platform.startswith("win"):
            os.startfile(str(target))  # type: ignore[attr-defined]
        else:
            subprocess.Popen([_opener_name(), str(target)], stdin=subprocess.DEVNULL,
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
        return {"path": str(target), "opened": True}

    return _audited("open_path", path, run)


def _command(argv: Any, cwd: Any) -> tuple[list[str], Path]:
    if isinstance(argv, (str, bytes)):
        raise ComputerAccessError("argv must be a list of arguments, not a command string")
    if not isinstance(argv, (list, tuple)) or not argv or not all(isinstance(a, str) and a for a in argv):
        raise ComputerAccessError("argv must be a non-empty list of non-empty strings")
    for arg in argv:
        bad = sorted(set(arg) & SHELL_METACHARACTERS)
        if bad or "\x00" in arg:
            raise ComputerAccessError(f"argument contains shell metacharacters {bad!r}; no shell is available")
    program = Path(argv[0]).name.lower()
    if program in SHELL_PROGRAMS or program.endswith((".bat", ".cmd")):
        raise ComputerAccessError(f"shell interpreters are not allowed: {argv[0]}")
    cfg = load_config()
    if cwd is None:
        if not cfg["allowed_roots"]:
            raise ComputerAccessError("no allowed roots are configured")
        cwd = cfg["allowed_roots"][0]
    workdir = guard(cwd, cfg)
    if not workdir.is_dir():
        raise ComputerAccessError(f"cwd is not a directory: {cwd}")
    return list(argv), workdir


def run_command(argv: Any, cwd: Any = None) -> dict[str, Any]:
    """Run argv (no shell) in cwd under the allowed roots with a timeout and capped output."""
    program = argv[0] if isinstance(argv, (list, tuple)) and argv else None

    def run() -> dict[str, Any]:
        args, workdir = _command(argv, cwd)
        runner = RUNNER or subprocess.run
        try:
            done = runner(args, cwd=str(workdir), shell=False, stdin=subprocess.DEVNULL, capture_output=True,
                          text=True, encoding="utf-8", errors="replace", timeout=COMMAND_TIMEOUT_SECONDS)
        except subprocess.TimeoutExpired:
            return {"path": str(workdir), "argv": args, "timed_out": True, "returncode": None}
        stdout, out_cut = _cap(done.stdout or "", MAX_OUTPUT_CHARS)
        stderr, err_cut = _cap(done.stderr or "", MAX_OUTPUT_CHARS)
        return {"path": str(workdir), "argv": args, "returncode": done.returncode, "timed_out": False,
                "stdout": stdout, "stderr": stderr, "output_truncated": out_cut or err_cut}

    return _audited("run_command", cwd, run, program=program)
